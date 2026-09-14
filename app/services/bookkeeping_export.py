"""Create a formatted workbook matching the bookkeeping file list."""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


def _text_cell(sheet, row, column, value):
    cell = sheet.cell(row, column)
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))
    # User-supplied filenames and names must remain text, including leading '='.
    cell.data_type = "s"
    return cell


def build_bookkeeping_workbook(groups, period):
    workbook = Workbook()
    workbook.remove(workbook.active)
    period_label = datetime.strptime(period, "%Y-%m").strftime("%B %Y")
    workbook.properties.title = f"Bookkeeping - {period_label}"
    workbook.properties.creator = "AXIODY"
    line = Side(style="thin", color="DCD6C8")
    for key, group in groups.items():
        sheet = workbook.create_sheet(group["label"])
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = "33513F"
        sheet.merge_cells("A1:C1")
        sheet["A1"] = group["label"]
        sheet["A1"].font = Font(name="Calibri", size=20, bold=True, color="16211C")
        count = len(group["files"])
        sheet["D1"] = f"{count} file{'s' if count != 1 else ''}"
        sheet["D1"].alignment = Alignment(horizontal="right", vertical="center")
        sheet["D1"].font = Font(name="Calibri", color="6B8577")
        sheet.row_dimensions[1].height = 34
        sheet.merge_cells("A2:D2")
        sheet["A2"] = period_label
        sheet["A2"].font = Font(name="Calibri", color="6B8577")

        for column, header in enumerate(["NAME", "SUBMITTED BY", "SIZE", "DATE"], 1):
            cell = sheet.cell(4, column, header)
            cell.font = Font(name="Calibri", bold=True, color="33513F")
            cell.fill = PatternFill("solid", fgColor="EFE9D8")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[4].height = 25

        for row_number, document in enumerate(group["files"], 5):
            values = [document["name"], f'{document["submittedBy"]}\n{document["submittedAt"]}',
                      document["size"], document["date"]]
            for column, value in enumerate(values, 1):
                cell = _text_cell(sheet, row_number, column, value)
                cell.font = Font(name="Calibri", size=11, bold=column == 1,
                                 color="16211C" if column <= 2 else "6B8577")
                cell.alignment = Alignment(vertical="center", wrap_text=True,
                                           horizontal="right" if column == 4 else "left")
                cell.border = Border(bottom=line)
            sheet.row_dimensions[row_number].height = 42

        # Keep a usable table on empty tabs, with one blank entry row.
        last_row = max(5, 4 + count)
        table = Table(displayName=f"Bookkeeping_{key}", ref=f"A4:D{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight14", showRowStripes=False)
        sheet.add_table(table)
        if not count:
            sheet["A7"] = f"No {group['label'].lower()} for this period."
            sheet["A7"].font = Font(name="Calibri", italic=True, color="6B8577")
        for column, width in {"A": 76, "B": 32, "C": 14, "D": 16}.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A5"
        sheet.print_title_rows = "1:4"
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = f"A1:D{max(7, last_row)}"

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output
